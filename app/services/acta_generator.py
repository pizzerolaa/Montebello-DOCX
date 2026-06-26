from __future__ import annotations

from dataclasses import replace
from functools import lru_cache
from pathlib import Path

from docx import Document
from docx.shared import Inches

from app.config import Settings
from app.models.equipment import Equipment
from app.models.project import Project
from app.services.equipment_format import normalize_serial
from app.services.excel_service import derived_work_done
from app.services.docx_report_appender import append_source_work_reports
from app.services.template_service import (
    ParagraphSpec,
    insert_paragraphs_at_marker,
    load_template,
    normalize_document_font,
    replace_placeholders,
)
from app.services.work_option_service import cold_room_work_for_equipment, minisplit_work_for_equipment, package_work_for_equipment


CONTENT_ROOT = Path(__file__).resolve().parents[1] / "document_content"


def generate_acta(project: Project, settings: Settings, output_path: Path) -> Path:
    """Generate ACTA DOCX from the configured Word template markers."""
    template_path = Path(__file__).resolve().parents[1] / "word_templates" / "acta_template.docx"
    document = load_template(template_path)
    _apply_document_layout(document)
    replace_placeholders(document, _placeholder_values(project))
    insert_paragraphs_at_marker(document, "[[SERVICE_SUMMARY]]", build_service_summary(project))
    insert_paragraphs_at_marker(document, "[[EQUIPMENT_SECTIONS]]", build_equipment_sections(project))
    normalize_document_font(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    append_source_work_reports(output_path, project)
    Document(output_path)
    return output_path


def build_service_summary(project: Project) -> list[ParagraphSpec]:
    counts = _type_counts(project)
    count_texts = []
    if counts["MINISPLIT"]:
        count_texts.append(_count_phrase(counts["MINISPLIT"], "EQUIPO TIPO MINISPLIT"))
    if counts["PACKAGE"]:
        count_texts.append(_count_phrase(counts["PACKAGE"], "EQUIPO TIPO PAQUETE"))
    if counts["COLD_ROOM"]:
        count_texts.append(_count_phrase(counts["COLD_ROOM"], "CAMARA FRIA", "CAMARAS FRIAS"))
    heading = "MANTENIMIENTO CORRECTIVO AL SISTEMA DE ENFRIAMIENTO DE AIRE ACONDICIONADO"
    if count_texts:
        heading = f"{heading} ({' Y '.join(count_texts)}), QUE INCLUYE:"
    paragraphs = [
        ParagraphSpec(
            heading,
            bold=True,
            alignment="justify",
            highlight="yellow",
            space_after_pt=12,
        )
    ]
    if counts["MINISPLIT"]:
        paragraphs.extend(_content_specs("minisplit", "service_summary.txt", _service_context(project, "MINISPLIT")))
    if counts["PACKAGE"]:
        paragraphs.extend(_content_specs("package", "service_summary.txt", _service_context(project, "PACKAGE")))
    if counts["COLD_ROOM"]:
        paragraphs.extend(_content_specs("cold_room", "service_summary.txt", _service_context(project, "COLD_ROOM")))
    return paragraphs


def build_equipment_sections(project: Project) -> list[ParagraphSpec]:
    paragraphs: list[ParagraphSpec] = []
    for index, equipment in enumerate(sorted(project.equipment, key=lambda item: item.sequence)):
        builder = ACTA_BUILDERS.get(equipment.equipment_type, build_unknown_acta_section)
        section = builder(equipment, index + 1)
        if index == 0:
            section[0] = replace(
                section[0],
                section_break_before=True,
                section_break_type="continuous",
                section_top_margin_inches=2.15,
            )
        paragraphs.extend(section)
    return paragraphs


def build_minisplit_acta_section(equipment: Equipment, number: int) -> list[ParagraphSpec]:
    paragraphs = _equipment_metadata(equipment, number)
    paragraphs.extend(_content_specs("minisplit", "acta_maintenance.txt"))
    work_done = minisplit_work_for_equipment(equipment) or derived_work_done(equipment)
    if work_done:
        paragraphs.extend(_content_specs("minisplit", "work_intro.txt"))
        paragraphs.append(ParagraphSpec(work_done, alignment="justify", space_after_pt=16))
    return paragraphs


def build_package_acta_section(equipment: Equipment, number: int) -> list[ParagraphSpec]:
    paragraphs = _equipment_metadata(equipment, number)
    paragraphs.extend(_content_specs("package", "acta_maintenance.txt"))
    work_items = package_work_for_equipment(equipment)
    work_done = derived_work_done(equipment)
    if work_items or work_done:
        paragraphs.append(
            ParagraphSpec(
                "PIEZAS SUSTITUIDAS Y CORRECCIONES REALIZADAS",
                bold=True,
                alignment="justify",
                space_before_pt=14,
                space_after_pt=12,
            )
        )
        paragraphs.extend(_content_specs("package", "work_intro.txt"))
        for line in work_items or work_done.splitlines():
            paragraphs.append(_piece_spec(line, number_sequence=f"work-{equipment.sequence}"))
        paragraphs.extend(_content_specs("package", "final_note.txt"))
    return paragraphs


def build_cold_room_acta_section(equipment: Equipment, number: int) -> list[ParagraphSpec]:
    paragraphs = _equipment_metadata(equipment, number)
    paragraphs.extend(_content_specs("cold_room", "acta_maintenance.txt"))
    work_items = cold_room_work_for_equipment(equipment)
    work_done = derived_work_done(equipment)
    if work_items or work_done:
        paragraphs.append(
            ParagraphSpec(
                "PIEZAS SUSTITUIDAS Y CORRECCIONES REALIZADAS",
                bold=True,
                alignment="justify",
                space_before_pt=14,
                space_after_pt=12,
            )
        )
        paragraphs.extend(_content_specs("cold_room", "work_intro.txt"))
        for line in work_items or work_done.splitlines():
            paragraphs.append(_piece_spec(line, number_sequence=f"work-{equipment.sequence}"))
        paragraphs.extend(_content_specs("cold_room", "final_note.txt"))
    return paragraphs


def build_unknown_acta_section(equipment: Equipment, number: int) -> list[ParagraphSpec]:
    paragraphs = _equipment_metadata(equipment, number)
    work_done = derived_work_done(equipment)
    if work_done:
        paragraphs.append(ParagraphSpec(work_done, alignment="justify"))
    return paragraphs


ACTA_BUILDERS = {
    "MINISPLIT": build_minisplit_acta_section,
    "PACKAGE": build_package_acta_section,
    "COLD_ROOM": build_cold_room_acta_section,
}


def _equipment_metadata(equipment: Equipment, number: int) -> list[ParagraphSpec]:
    title = f"{number}.  ZONA DE {equipment.zone or 'PENDIENTE'}"
    paragraphs = [
        ParagraphSpec(
            title,
            bold=True,
            alignment="justify",
            space_before_pt=6,
            space_after_pt=8,
        ),
        ParagraphSpec(f"MARCA: **{equipment.brand or ''}**", alignment="justify"),
    ]
    if equipment.capacity:
        paragraphs.append(ParagraphSpec(f"CAPACIDAD: **{equipment.capacity}**", alignment="justify"))
    paragraphs.append(
        ParagraphSpec(
            f"NUMERO DE SERIE: **{normalize_serial(equipment.serial)}**",
            alignment="justify",
            space_after_pt=18,
        )
    )
    return paragraphs


def _placeholder_values(project: Project) -> dict[str, str]:
    deliverer = next((item for item in project.signatures if item.role == "DELIVERER"), None)
    return {
        "LOCATION": project.location or "",
        "DATE": project.service_date_raw or "",
        "ORDER_NUMBER": project.order_number or "",
        "CENTER_NAME": project.center_name or "",
        "DELIVERER_NAME": deliverer.name if deliverer else "",
        "DELIVERER_POSITION": deliverer.position if deliverer else "",
    }


def _type_counts(project: Project) -> dict[str, int]:
    return {
        "MINISPLIT": sum(1 for equipment in project.equipment if equipment.equipment_type == "MINISPLIT"),
        "PACKAGE": sum(1 for equipment in project.equipment if equipment.equipment_type == "PACKAGE"),
        "COLD_ROOM": sum(1 for equipment in project.equipment if equipment.equipment_type == "COLD_ROOM"),
    }


def _count_phrase(count: int, label: str, plural_label: str | None = None) -> str:
    noun = label if count == 1 else plural_label or label.replace("EQUIPO", "EQUIPOS", 1)
    return f"{count:02d} {noun}"


def _content_specs(equipment_folder: str, filename: str, context: dict[str, str] | None = None) -> list[ParagraphSpec]:
    return [
        ParagraphSpec(_format_content(text, context), alignment="justify", space_after_pt=10)
        for text in _load_content(equipment_folder, filename)
    ]


def _piece_spec(line: str, *, number_sequence: str = "work-items") -> ParagraphSpec:
    if ":" not in line:
        return ParagraphSpec(
            line,
            alignment="justify",
            style="List Bullet",
            number_sequence=number_sequence,
            space_after_pt=2,
        )
    title, detail = line.split(":", 1)
    return ParagraphSpec(
        f"**{title.strip()}:**{detail}",
        alignment="justify",
        style="List Bullet",
        number_sequence=number_sequence,
        space_after_pt=2,
    )


def _service_context(project: Project, equipment_type: str) -> dict[str, str]:
    equipment = [item for item in project.equipment if item.equipment_type == equipment_type]
    count = len(equipment)
    equipment_noun = "EQUIPO" if count == 1 else "EQUIPOS"
    if equipment_type == "COLD_ROOM":
        equipment_noun = "CAMARA FRIA" if count == 1 else "CAMARAS FRIAS"
    return {
        "COUNT": str(count),
        "COUNT_02": f"{count:02d}",
        "EQUIPMENT_NOUN": equipment_noun,
        "CAPACITY_CLAUSE": _capacity_clause(equipment, equipment_type),
    }   


def _capacity_clause(equipment: list[Equipment], equipment_type: str) -> str:
    capacities = []
    for item in equipment:
        capacity = (item.capacity or "").strip().rstrip(".")
        if capacity and capacity.upper() not in {value.upper() for value in capacities}:
            capacities.append(capacity)
    if not capacities:
        return ""
    if len(capacities) == 1:
        if equipment_type == "COLD_ROOM":
            return f" CON CAPACIDAD DE {capacities[0]}"
        if equipment_type == "PACKAGE":
            return f" CON CAPACIDAD DE {capacities[0]}"
        return f" DE {capacities[0]}"
    if equipment_type == "COLD_ROOM":
        return f" CON CAPACIDADES DE {', '.join(capacities)}"
    if equipment_type == "PACKAGE":
        return f" CON CAPACIDADES DE {', '.join(capacities)}"
    return f" DE {', '.join(capacities)}"


def _format_content(text: str, context: dict[str, str] | None) -> str:
    if not context:
        return text
    return text.format_map(_SafeFormatMap(context))


class _SafeFormatMap(dict[str, str]):
    def __missing__(self, key: str) -> str:
        return "{" + key + "}"


def _apply_document_layout(document) -> None:
    for section in document.sections:
        section.bottom_margin = Inches(1.25)


@lru_cache
def _load_content(equipment_folder: str, filename: str) -> tuple[str, ...]:
    path = CONTENT_ROOT / equipment_folder / filename
    if not path.exists():
        return ()
    text = path.read_text(encoding="utf-8")
    return tuple(part.strip() for part in text.split("\n\n") if part.strip())
